"""Alianzas y encuestas 2027 que ve todo el equipo (no solo quien las captura en su pantalla) → meta.esc.

Se llenan en un Excel (BASE/ESCENARIOS/escenarios_2027.xlsx) con dos hojas:
  Alianzas:  estado, cargo, partidos que van juntos, tipo (coalición o candidatura común), fecha, fuente, activo.
  Encuestas: estado, cargo, casa encuestadora, fecha de levantamiento, muestra, margen de error, % por partido, activo.
Varias encuestas del mismo estado y cargo se promedian con más peso a las recientes (la mitad del peso cada 30 días)
y a las de muestra más grande (hasta 3,000), que es la forma usual de promediar encuestas. El margen que se muestra
es el promedio ponderado de los márgenes publicados (conservador).

uso: python pipeline/lab.py escenarios            (lee el Excel y actualiza privado/<estado>.json)
     python pipeline/lab.py escenarios --plantilla
"""
import datetime as dt
import json
import sys
from pathlib import Path

from comun import BASE, PLANO, estado

XLS = BASE / 'ESCENARIOS' / 'escenarios_2027.xlsx'
CARGOS = {'GOB': 'Gubernatura', 'DIPL': 'Diputaciones locales', 'AYU': 'Ayuntamientos', 'DIPF': 'Diputaciones federales'}
PARTIDOS = ['MORENA', 'PAN', 'PRI', 'PVEM', 'PT', 'MC', 'OTROS']
MEDIA_VIDA = 30


def plantilla(ruta=XLS):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if ruta.exists():
        raise SystemExit(f'Ya existe {ruta}. No lo sobrescribo.')
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    def hoja(ws, cols, ejemplos, anchos):
        ws.append(cols)
        for r in ejemplos:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1C7F79')
        for j, w in enumerate(anchos):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = w
        ws.freeze_panes = 'A2'
    a = wb.active; a.title = 'Alianzas'
    hoja(a, ['ESTADO', 'CARGO', 'PARTIDOS', 'TIPO', 'FECHA', 'FUENTE', 'ACTIVO'],
         [['camp', 'GOB', 'MORENA, PT, PVEM', 'COALICION', '2027-01-15', 'Ejemplo: convenio registrado ante el instituto electoral', 'NO'],
          ['camp', 'GOB', 'PAN, PRI', 'COALICION', '2027-01-15', 'Ejemplo', 'NO']], (10, 8, 28, 16, 12, 50, 8))
    dv = DataValidation(type='list', formula1='"GOB,DIPL,AYU,DIPF"'); a.add_data_validation(dv); dv.add('B2:B500')
    dv = DataValidation(type='list', formula1='"COALICION,CANDIDATURA COMUN"'); a.add_data_validation(dv); dv.add('D2:D500')
    dv = DataValidation(type='list', formula1='"SI,NO"'); a.add_data_validation(dv); dv.add('G2:G500')
    e = wb.create_sheet('Encuestas')
    hoja(e, ['ESTADO', 'CARGO', 'CASA ENCUESTADORA', 'FECHA', 'MUESTRA', 'MARGEN', *PARTIDOS, 'ACTIVO'],
         [['camp', 'GOB', 'Ejemplo Encuestas SA', '2027-02-10', 1000, 3.1, 45, 18, 12, 6, 4, 10, 5, 'NO']], (10, 8, 26, 12, 10, 9) + (9,) * len(PARTIDOS) + (8,))
    dv = DataValidation(type='list', formula1='"GOB,DIPL,AYU,DIPF"'); e.add_data_validation(dv); dv.add('B2:B500')
    dv = DataValidation(type='list', formula1='"SI,NO"'); e.add_data_validation(dv); dv.add(f'{openpyxl.utils.get_column_letter(7 + len(PARTIDOS))}2:{openpyxl.utils.get_column_letter(7 + len(PARTIDOS))}500')
    i = wb.create_sheet('Instrucciones')
    for t in ['Cómo llenar este archivo', '',
              'ESTADO: clave (camp, chis, pue…) o nombre del estado.',
              'CARGO: GOB gubernatura · DIPL diputaciones locales · AYU ayuntamientos · DIPF diputaciones federales.',
              'Hoja Alianzas · PARTIDOS: siglas separadas por coma (MORENA, PT, PVEM). Una fila por alianza. Los partidos que no estén en ninguna alianza van solos.',
              'Hoja Alianzas · TIPO: COALICION o CANDIDATURA COMUN.',
              'Hoja Encuestas · porcentajes de la intención de voto efectiva (sin indecisos), por partido. OTROS = todo lo demás. Deben sumar cerca de 100.',
              'Hoja Encuestas · MUESTRA: número de entrevistas. MARGEN: margen de error publicado, en puntos.',
              'Varias encuestas del mismo estado y cargo se promedian: pesan más las recientes y las de muestra más grande.',
              'FECHA: AAAA-MM-DD (en encuestas, el último día de levantamiento).',
              'ACTIVO: SI para que cuente; NO para dejarla guardada sin usarla.',
              '',
              'Después de editar: avísale a Claude o corre  python pipeline/lab.py escenarios  y luego  python pipeline/lab.py sitio.',
              'Lo que se captura aquí lo ve todo el equipo. En pantalla cada quien puede probar otras alianzas o encuestas sin cambiar esto.']:
        i.append([t])
    i['A1'].font = Font(bold=True, size=13); i.column_dimensions['A'].width = 130
    wb.save(ruta)
    print(f'[escenarios] plantilla en {ruta}')


def _fecha(x):
    if isinstance(x, (dt.date, dt.datetime)):
        return x if isinstance(x, dt.date) and not isinstance(x, dt.datetime) else x.date()
    try:
        return dt.date.fromisoformat(str(x).strip()[:10])
    except ValueError:
        return None


def _f(x):
    try:
        return float(str(x).replace('%', '').replace(',', '').strip())
    except (TypeError, ValueError):
        return None


def leer(ruta=XLS):
    import openpyxl
    if not ruta.exists():
        return {}, []
    wb = openpyxl.load_workbook(ruta, data_only=True); out = {}; avisos = []
    def filas(nombre):
        if nombre not in wb.sheetnames:
            return [], []
        R = list(wb[nombre].iter_rows(values_only=True))
        return [str(h or '').strip().upper() for h in R[0]], [r for r in R[1:] if r and any(v not in (None, '') for v in r)]
    H, R = filas('Alianzas')
    for n, r in enumerate(R, 2):
        g = dict(zip(H, r))
        if str(g.get('ACTIVO') or '').strip().upper() != 'SI':
            continue
        try:
            cab = estado(g['ESTADO'])['clave']
        except SystemExit:
            avisos.append(f'Alianzas fila {n}: estado no reconocido'); continue
        c = str(g.get('CARGO') or '').strip().upper()
        if c not in CARGOS:
            avisos.append(f'Alianzas fila {n}: cargo {c!r} no válido'); continue
        ps = [p.strip().upper() for p in str(g.get('PARTIDOS') or '').replace(';', ',').split(',') if p.strip()]
        if len(ps) < 2:
            avisos.append(f'Alianzas fila {n}: una alianza necesita al menos dos partidos'); continue
        f = _fecha(g.get('FECHA'))
        out.setdefault(cab, {'al': {}, 'enc': {}})['al'].setdefault(c, []).append(
            {'p': ps, 'tipo': str(g.get('TIPO') or '').strip().lower(), 'fecha': f.isoformat() if f else None, 'fuente': str(g.get('FUENTE') or '').strip()})
    H, R = filas('Encuestas'); hoy = dt.date.today(); grupos = {}
    for n, r in enumerate(R, 2):
        g = dict(zip(H, r))
        if str(g.get('ACTIVO') or '').strip().upper() != 'SI':
            continue
        try:
            cab = estado(g['ESTADO'])['clave']
        except SystemExit:
            avisos.append(f'Encuestas fila {n}: estado no reconocido'); continue
        c = str(g.get('CARGO') or '').strip().upper(); f = _fecha(g.get('FECHA')); m = _f(g.get('MUESTRA')) or 800; e = _f(g.get('MARGEN')) or 3.5
        v = {p: _f(g.get(p)) for p in PARTIDOS if _f(g.get(p)) is not None}
        if c not in CARGOS or not f or not v:
            avisos.append(f'Encuestas fila {n}: falta cargo, fecha o porcentajes'); continue
        s = sum(v.values())
        if not 90 <= s <= 110:
            avisos.append(f'Encuestas fila {n}: los porcentajes suman {s:.0f}; se reescalan a 100')
        v = {p: 100 * x / s for p, x in v.items()}
        w = min(m, 3000) / 1000 * 0.5 ** (max(0, (hoy - f).days) / MEDIA_VIDA)
        grupos.setdefault((cab, c), []).append({'casa': str(g.get('CASA ENCUESTADORA') or '').strip(), 'fecha': f.isoformat(), 'n': int(m), 'moe': e, 'v': v, 'w': w})
    for (cab, c), L in grupos.items():
        W = sum(x['w'] for x in L) or 1
        prom = {p: round(sum(x['w'] * x['v'].get(p, 0) for x in L) / W, 1) for p in PARTIDOS if any(p in x['v'] for x in L)}
        prom.pop('OTROS', None)  # "otros" se reparte entre el resto de partidos, como en la captura en pantalla
        out.setdefault(cab, {'al': {}, 'enc': {}})['enc'][c] = {'v': prom, 'moe': round(sum(x['w'] * x['moe'] for x in L) / W, 1),
                                                                 'n': len(L), 'ult': max(x['fecha'] for x in L), 'lista': [{k: x[k] for k in ('casa', 'fecha', 'n', 'moe')} for x in L]}
    return out, avisos


def generar(ruta=XLS):
    out, avisos = leer(ruta)
    for a in avisos:
        print('   aviso ·', a)
    for f in sorted(PLANO.glob('*.json')):
        if f.stem == 'claves':
            continue
        d = json.loads(f.read_text(encoding='utf-8')); cab = d['meta']['cab']
        nuevo = out.get(cab)
        if (d['meta'].get('esc') or None) == nuevo:
            continue
        if nuevo:
            d['meta']['esc'] = nuevo
        else:
            d['meta'].pop('esc', None)
        f.write_text(json.dumps(d, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'[escenarios] {sum(len(v["al"]) for v in out.values())} cargos con alianza y {sum(len(v["enc"]) for v in out.values())} con encuestas en {len(out)} estados')


if __name__ == '__main__':
    if '--plantilla' in sys.argv:
        plantilla(Path(sys.argv[sys.argv.index('--plantilla') + 1]) if len(sys.argv) > sys.argv.index('--plantilla') + 1 else XLS)
    else:
        generar()
