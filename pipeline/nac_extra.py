"""Paso 4c. Datos complementarios de la vista nacional del inicio (requiere openpyxl).

Genera data/inicio/nac_extra.json con:
  leg   integración de la Cámara (2015, 2018, 2021, 2024) y del Senado (2012, 2018, 2024) con nombres
  mun   partido que gobierna cada municipio (de pipeline/insumos/ayuntamientos_edad.json, que genera pipeline/extra_inicio.py)
  pad   lista nominal por estado: total, mujeres y de 18 a 29 años; y municipios que reúnen el 70% del electorado
Fuentes: INE, catálogos de integración de cada proceso federal; DERFE, padrón y lista nominal por sección.

uso: python pipeline/nac_extra.py
"""
import collections, csv, json, re
import openpyxl
from comun import BASE, CFG, PIPE, REPO, ruta

MAIN = ['MORENA', 'PAN', 'PRI', 'PRD', 'PVEM', 'PT', 'MC']
ALIAS = {'MOVIMIENTO': 'MC', 'MOV': 'MC', 'CIUDADANO': 'MC', 'PANAL': 'NA', 'NVA': 'NA', 'ES': 'PES'}
CARGO = {'DIP_MR': 'MR', 'DIP_RP': 'RP', 'SEN_MR': 'MR', 'SEN_RP': 'LN', 'SEN_PM': 'PM'}


def lider(actor):
    ts = [ALIAS.get(t, t) for t in re.split(r'[_\s]+', str(actor).upper().replace('COM_MEXICO', 'PRI_PVEM')) if t and t not in ('COM', 'C', 'CC', 'ALIANZA')]
    if any('IND' in t for t in ts):
        return 'IND'
    return next((p for p in MAIN if p in ts), ts[0] if ts else 'OTROS')


def legislaturas():
    cat = ruta('federal') / '_CATALOGOS'; L = {}
    for y in (2012, 2015, 2018, 2021):
        ws = openpyxl.load_workbook(cat / str(y) / f'Integracion_PEF_{y}.xlsx', read_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True); ix = {h: i for i, h in enumerate(next(it))}; out = []
        for r in it:
            c = r[ix['CLAVE_CARGO_POLITICO']]
            if c not in CARGO or (y == 2021 and c.startswith('SEN')) or (y == 2012 and c.startswith('DIP')):
                continue
            out.append([c[:3], lider(r[ix['ACTOR_POLITICO']]), (r[ix['CANDIDATO']] or '').title(), int(r[ix['ID_ESTADO']] or 0), int(r[ix['ID_DISTRITO_FEDERAL']] or 0), int(r[ix['CIRCUNSCRIPCION']] or 0), CARGO[c]])
        por = collections.defaultdict(list)
        for x in out:
            if x[0] == 'SEN' and x[6] == 'MR':
                por[x[3]].append(x)
        for xs in por.values():
            if len(xs) == 3:
                xs[2][6] = 'PM'
        L[y] = out
    out = []
    for r in csv.DictReader(open(cat / '2024' / 'INTEGRACION_CARGOS_PEF_2024.csv', encoding='utf-8-sig')):
        t = r['TIPO_DE_CANDIDATURA']
        if t == 'PRE':
            continue
        p = 'IND' if 'INDEP' in r['PARTIDO_POLITICO'] else r['PARTIDO_POLITICO']
        h = 'PM' if t == 'SEN_MR' and r['NUMERO_LISTA'] == '3' else CARGO[t]
        out.append([t[:3], p, (r['PERSONA_CANDIDATA'] or '').title(), int(r['ID_ESTADO'] or 0), int(r['ID_DISTRITO_FEDERAL'] or 0), int(r['CIRCUNSCRIPCION'] or 0), h])
    L[2024] = out
    return L


def padron():
    ws = openpyxl.load_workbook(BASE / CFG['rutas']['padron'].replace('pdln_edms_re_', 'pdln_edms_re_'), read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True); H = [str(h).replace('\n', ' ') for h in next(it)]; ix = {h: i for i, h in enumerate(H)}
    jov = ['18', '19', '20_24', '25_29']; E = collections.defaultdict(lambda: [0, 0, 0]); MU = collections.Counter()
    for r in it:
        if r[0] is None or not str(r[0]).isdigit() or not r[4]:
            continue
        ln = r[ix['LISTA NOMINAL']]; m = r[ix['LISTA MUJERES']]; j = sum(r[ix[f'LISTA_{a}_{s}']] for a in jov for s in ('HOMBRES', 'MUJERES'))
        for k in ('0', str(r[0])):
            a = E[k]; a[0] += ln; a[1] += m; a[2] += j
        MU[f'{r[0]}_{r[3]}'] += ln
    con = {}
    for k in E:
        ms = sorted(((v, m) for m, v in MU.items() if k == '0' or m.split('_')[0] == k), reverse=True); tot = sum(v for v, _ in ms); acc = 0; n = 0
        for v, _ in ms:
            acc += v; n += 1
            if acc >= .7 * tot:
                break
        con[k] = [n, len(ms)]
    return {k: v + con[k] for k, v in E.items()}


def main():
    ex = json.load(open(PIPE / 'insumos' / 'ayuntamientos_edad.json'))
    out = {'leg': {str(k): v for k, v in legislaturas().items()}, 'mun': ex['mun'], 'pad': padron()}
    (REPO / 'data' / 'inicio' / 'nac_extra.json').write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print('data/inicio/nac_extra.json listo')


if __name__ == '__main__':
    main()
